"""Create only synthetic Northstar knowledge content in the chosen Confluence space.

Credentials are read from CONFLUENCE_EMAIL and CONFLUENCE_API_TOKEN. Re-running
reuses pages by exact title, and never edits pre-existing non-demo content.
"""

import argparse
import io
import json
import os
from pathlib import Path

import requests
from docx import Document
from reportlab.pdfgen import canvas

PERMISSION_NOTICE = (
    "<p><strong>Demo permission note:</strong> Page restrictions could not be applied "
    "with this account. Restriction scenarios described below are illustrative only. "
    "These pages retain the existing space access rules. The integration extracts "
    "the actual permission metadata; the shared Coveo source does not enforce it.</p>"
)

PAGES = [
    (
        "home",
        None,
        "Knowledge Hub",
        "Knowledge management",
        "none",
        """
        <h2>Northstar Cloud knowledge hub</h2>
        <p>Synthetic sales-demo content: all products, customers, incidents,
        and policies are fictional.</p>
        <p>Northstar Cloud provides managed analytics for enterprise support teams.
        This knowledge hub connects technical guides, troubleshooting articles,
        operating procedures, and attached handbooks.</p>
        <h3>Search tasks</h3><ul><li>How do I fix a SAML audience mismatch?</li>
        <li>When should an agent escalate a priority-one incident?</li>
        <li>What is the recovery point objective for Atlas?</li></ul>
        <p>Permission metadata demonstrates unrestricted-in-space, direct-user, direct-group,
        and inherited restrictions. The demo Coveo source is shared;
        it does not enforce those rules.</p>""",
    ),
    (
        "sso",
        "home",
        "Resolve SAML audience mismatch",
        "Support",
        "none",
        """
        <h2>KB-1042: SSO login fails with audience mismatch</h2>
        <p>Applies to Northstar Atlas Enterprise 4.2. Owner: Identity Support.</p>
        <p>Symptom: users see AUTH-SAML-041 after an identity-provider migration.</p>
        <h3>Resolution</h3><ol><li>Export the service-provider metadata
        from Admin &gt; Identity.</li>
        <li>Set the IdP audience URI to urn:northstar:atlas:production.</li>
        <li>Verify the assertion consumer URL and certificate fingerprint.</li>
        <li>Allow a five-minute clock skew, then test using an incognito browser session.</li></ol>
        <p>Do not disable signature verification.
        Escalate if the audience matches and failures persist.</p>
        <p>Verification phrase: northstar-orchid-saml.
        Related attachment: identity troubleshooting checklist.</p>""",
    ),
    (
        "api",
        "home",
        "Atlas API ingestion technical guide",
        "Engineering",
        "group",
        """
        <h2>Atlas ingestion API: retries and idempotency</h2>
        <p>For connector engineers integrating a new source into Northstar Atlas.</p>
        <p>Use POST /v2/events with an Idempotency-Key unique to the source event.
        Retain the key for
        24 hours. Replaying the same key returns the original event receipt without duplication.</p>
        <p>Batch at most 500 events and 5 MB. A 429 response includes Retry-After; apply exponential
        backoff with jitter. Retry 502/503/504, but fix a 400 payload before resubmission.</p>
        <h3>Operational checks</h3><p>Monitor accepted_events, rejected_events,
        and ingestion_lag_seconds.
        A lag above 120 seconds for ten minutes creates an engineering alert.</p>""",
    ),
    (
        "sop",
        "home",
        "Priority-one incident response SOP",
        "Service Operations",
        "none",
        """
        <h2>SOP-017: Priority-one incident response</h2>
        <p>A priority-one incident means production analytics
        are unavailable for multiple customers.</p>
        <ol><li>Within five minutes, assign an incident commander
        and open the incident channel.</li>
        <li>Within fifteen minutes, publish an initial status update
        and record customer impact.</li>
        <li>Check the release timeline, service health,
        and ingestion queue before rolling back.</li>
        <li>Post updates every thirty minutes; preserve logs and a timeline of decisions.</li>
        <li>After recovery, complete a blameless review within two business days.</li></ol>
        <p>Verification phrase: northstar-cobalt-escalation.
        See the attached incident-response playbook.</p>""",
    ),
    (
        "restricted",
        "home",
        "Restricted engineering operations",
        "Engineering",
        "user",
        """
        <h2>Restricted engineering operations</h2>
        <p>This synthetic collection represents internal operational knowledge
        limited to an owner.</p>
        <p>Child pages inherit this page's read restriction.
        Attachments inherit their containing page's
        permissions. The integration must carry each restriction layer, not merge them with OR.</p>
        <p>No real secrets, credentials, customer data,
        or production procedures appear here.</p>""",
    ),
    (
        "recovery",
        "restricted",
        "Atlas disaster recovery runbook",
        "Engineering",
        "none",
        """
        <h2>Atlas disaster recovery: regional failover</h2>
        <p>Fictional objectives: recovery time objective 60 minutes;
        recovery point objective 15 minutes.</p>
        <ol><li>Confirm a regional outage with the incident commander.</li>
        <li>Verify the most recent replica checkpoint
        and pause ingestion to prevent split-brain writes.</li>
        <li>Promote the standby region, run integrity checks, and then shift traffic gradually.</li>
        <li>Resume ingestion and reconcile events using source event identifiers.</li></ol>
        <p>Verification phrase: northstar-amber-failover. This page has no direct read restriction;
        its parent restriction is essential to interpreting effective access.</p>""",
    ),
    (
        "onboarding",
        "home",
        "Support agent onboarding wiki",
        "Customer Success",
        "none",
        """
        <h2>New support agent: first-week checklist</h2>
        <p>Day one: complete security training and request the Support Analyst role.</p>
        <p>Day two: shadow two customer cases and learn the difference between a knowledge article,
        a technical guide, and an operational runbook.</p>
        <p>Day three: practice finding SAML troubleshooting
        and incident escalation guidance using search.</p>
        <p>Before replying to a customer, verify product version,
        article review date, and audience.</p>
        <p>Never copy restricted engineering material into a public support response.</p>""",
    ),
    (
        "relevance",
        "home",
        "Knowledge quality and search relevance guide",
        "Knowledge Management",
        "group",
        """
        <h2>Improve knowledge discoverability</h2>
        <p>Use a task-oriented title, an explicit product version,
        a named owner, and a review date.</p>
        <p>Measure zero-result queries, click-through rate,
        and successful self-service journeys weekly.</p>
        <p>When a query repeatedly fails, inspect synonyms
        and missing content before changing ranking.</p>
        <p>Examples: map SSO to single sign-on; connect P1 to priority-one incident; label recovery
        documents with both RTO and recovery time objective.</p>
        <p>A unified index should find answers in page bodies and PDF/Word attachments without
        requiring employees to know which repository contains the answer.</p>""",
    ),
]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--site", required=True)
    parser.add_argument("--space-id", required=True)
    parser.add_argument("--group-id", required=True)
    parser.add_argument("--manifest", required=True)
    parser.add_argument(
        "--skip-restrictions",
        action="store_true",
        help="Create content without changing page restrictions when the API denies it",
    )
    args = parser.parse_args()
    base = args.site.rstrip("/").removesuffix("/wiki")
    session = requests.Session()
    session.auth = (os.environ["CONFLUENCE_EMAIL"], os.environ["CONFLUENCE_API_TOKEN"])

    def call(method, path, **kwargs):
        response = session.request(method, base + path, timeout=60, **kwargs)
        response.raise_for_status()
        return response.json()

    owner = call("GET", "/wiki/rest/api/user/current")["accountId"]
    manifest = {"site_url": base, "space_id": args.space_id, "pages": {}, "attachments": []}
    output = Path(args.manifest)
    output.parent.mkdir(parents=True, exist_ok=True)

    def save():
        output.write_text(json.dumps(manifest, indent=2))

    for key, parent, title, department, restriction, body in PAGES:
        title = "[Coveo Demo] " + title
        matches = call(
            "GET", "/wiki/api/v2/pages", params={"space-id": args.space_id, "title": title}
        )["results"]
        if matches:
            page = matches[0]
        else:
            payload = {
                "spaceId": args.space_id,
                "status": "current",
                "title": title,
                "body": {
                    "representation": "storage",
                    "value": (PERMISSION_NOTICE if args.skip_restrictions else "") + body,
                },
            }
            if parent:
                payload["parentId"] = manifest["pages"][parent]["id"]
            page = call("POST", "/wiki/api/v2/pages", json=payload)
            page_id = page["id"]
            call(
                "POST",
                f"/wiki/rest/api/content/{page_id}/label",
                json=[
                    {"prefix": "global", "name": "coveo-demo"},
                    {"prefix": "global", "name": department.lower().replace(" ", "-")},
                ],
            )
            call(
                "POST",
                f"/wiki/api/v2/pages/{page_id}/properties",
                json={
                    "key": "demo_knowledge",
                    "value": {
                        "department": department,
                        "product": "Northstar Atlas",
                        "synthetic": True,
                        "review_date": "2026-10-21",
                    },
                },
            )
            if restriction != "none" and not args.skip_restrictions:
                access = {"user": [{"type": "known", "accountId": owner}], "group": []}
                if restriction == "group":
                    access["group"] = [{"type": "group", "id": args.group_id}]
                call(
                    "PUT",
                    f"/wiki/rest/api/content/{page_id}/restriction",
                    json=[{"operation": "read", "restrictions": access}],
                )
        manifest["pages"][key] = {
            "id": page["id"],
            "title": title,
            "intended_restriction_scenario": restriction,
            "restrictions_skipped": args.skip_restrictions,
        }
        save()
        print("Demo page ready:", key, page["id"])

    pdf = io.BytesIO()
    document = canvas.Canvas(pdf)
    for index, line in enumerate(
        [
            "Northstar Cloud - Synthetic Incident Response Playbook",
            "Verification phrase: northstar-sapphire-pdf",
            "P1 acknowledgement: 5 minutes. Initial customer update: 15 minutes.",
            "Appoint an incident commander; preserve a decision timeline.",
            "Publish updates every 30 minutes and review within two business days.",
        ]
    ):
        document.drawString(45, 780 - index * 24, line)
    document.save()
    docx = io.BytesIO()
    document = Document()
    document.add_heading("Northstar Atlas Recovery Checklist", 0)
    document.add_paragraph("Synthetic demo. Verification phrase: northstar-jade-docx.")
    document.add_paragraph(
        "RTO: 60 minutes. RPO: 15 minutes. Confirm replica checkpoints before failover."
    )
    document.add_paragraph(
        "The incident commander approves promotion; validate integrity before restoring traffic."
    )
    document.save(docx)
    attachments = [
        ("sop", "northstar-incident-playbook.pdf", "application/pdf", pdf.getvalue()),
        (
            "recovery",
            "northstar-recovery-checklist.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            docx.getvalue(),
        ),
        (
            "sso",
            "northstar-identity-checklist.txt",
            "text/plain",
            b"Synthetic Northstar SSO checklist. Verification phrase: northstar-topaz-text.\n"
            b"Verify audience URI, signing certificate and assertion consumer URL. "
            b"Never disable signatures.\n",
        ),
    ]
    for key, filename, mime, data in attachments:
        page_id = manifest["pages"][key]["id"]
        existing = call(
            "GET", f"/wiki/api/v2/pages/{page_id}/attachments", params={"filename": filename}
        )["results"]
        if existing:
            attachment = existing[0]
        else:
            attachment = call(
                "POST",
                f"/wiki/rest/api/content/{page_id}/child/attachment",
                headers={"X-Atlassian-Token": "nocheck"},
                files={"file": (filename, data, mime)},
            )["results"][0]
        manifest["attachments"].append(
            {"id": attachment["id"], "pageId": page_id, "title": filename}
        )
        save()
        print("Demo attachment ready:", filename, attachment["id"])


if __name__ == "__main__":
    main()
